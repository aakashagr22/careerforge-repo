import os
import zipfile
import copy
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def test_proper_order():
    template_path = r'c:\Users\hp\.gemini\antigravity\scratch\careerforge-repo\7th Sem Summer Training Report Format.docx'
    with zipfile.ZipFile(template_path) as z:
        img_data = z.read('word/media/image1.jpeg')
        with open('temp_image1.jpeg', 'wb') as f:
            f.write(img_data)

    doc_tpl = Document(template_path)
    doc = Document()

    # Register image
    pic_p = doc.add_paragraph()
    run = pic_p.add_run()
    inline_pic = run.add_picture('temp_image1.jpeg')
    blip = inline_pic._inline.xpath('.//a:blip')[0]
    img_rid = blip.get(qn('r:embed'))
    doc._body._element.remove(pic_p._p)

    # Clean initial empty paragraph from doc
    initial_p = doc.paragraphs[0]._p

    # Add page 2 paragraphs from template IN ORDER into section 0
    for i in range(41, 75):
        tpl_p = doc_tpl.paragraphs[i]
        cloned_p = copy.deepcopy(tpl_p._p)
        for br in cloned_p.xpath('.//w:br[@w:type="page"]'):
            br.getparent().remove(br)
        for br in cloned_p.xpath('.//w:lastRenderedPageBreak'):
            br.getparent().remove(br)
        for b in cloned_p.xpath('.//a:blip'):
            b.set(qn('r:embed'), img_rid)
        doc._body._element.append(cloned_p)

    # Remove the initial empty paragraph now that cloned paragraphs are inserted
    doc._body._element.remove(initial_p)

    # Now add section 1 (Preliminaries)
    sec_prelim = doc.add_section()
    doc.add_paragraph("Preliminary Page 1: Vision and Mission")

    print("Paragraphs in test doc:")
    for idx, p in enumerate(doc.paragraphs[:15]):
        print(f"{idx}: {repr(p.text)}")

test_proper_order()
